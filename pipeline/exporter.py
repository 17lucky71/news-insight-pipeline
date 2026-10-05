"""데이터 내보내기 모듈 (export 명령).

지원 형식: CSV, Excel(.xlsx), JSONL
필터: --status (cleaned / summarized), --category
"""
import csv
import json
import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)

# (DB 컬럼, 내보낼 때 쓰는 이름)
COLUMNS = [
    ("id", "ID"), ("category", "카테고리"), ("title", "제목"), ("published_at", "발행일시"),
    ("source", "출처"), ("status", "상태"), ("summary", "AI 요약"), ("sentiment", "감성"),
    ("url", "URL"), ("content", "본문"), ("collected_at", "수집일시"),
]
EXTENSIONS = {"csv": "csv", "excel": "xlsx", "jsonl": "jsonl"}


def export_csv(rows: list[dict], path: Path) -> None:
    # utf-8-sig: 엑셀에서 CSV 를 바로 열어도 한글이 깨지지 않도록 BOM 추가
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([name for _, name in COLUMNS])
        for r in rows:
            writer.writerow([r.get(col) or "" for col, _ in COLUMNS])


def export_jsonl(rows: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps({col: r.get(col) for col, _ in COLUMNS}, ensure_ascii=False) + "\n")


def export_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    ws.append([name for _, name in COLUMNS])
    for r in rows:
        ws.append([r.get(col) or "" for col, _ in COLUMNS])

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = {"ID": 6, "카테고리": 10, "제목": 45, "발행일시": 20, "출처": 10, "상태": 12,
              "AI 요약": 60, "감성": 8, "URL": 40, "본문": 60, "수집일시": 20}
    for i, (_, name) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(name, 15)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in ("C", "G"))
    ws.freeze_panes = "A2"  # 머리글 고정
    wb.save(path)


EXPORTERS = {"csv": export_csv, "excel": export_excel, "jsonl": export_jsonl}


def run_export(args, config: dict, storage) -> None:
    rows = storage.get_articles(status=args.status, category=args.category)
    rows.sort(key=lambda r: r["id"])
    if not rows:
        logger.warning("조건에 맞는 기사가 없습니다 (status=%s, category=%s).", args.status, args.category)
        return

    export_dir = Path(config.get("output", {}).get("export_dir", "output/exports"))
    export_dir.mkdir(parents=True, exist_ok=True)
    suffix = f"_{args.status}" if args.status else ""
    path = export_dir / f"news{suffix}_{datetime.now():%Y%m%d_%H%M%S}.{EXTENSIONS[args.format]}"

    EXPORTERS[args.format](rows, path)
    logger.info("내보내기 완료: %d건 → %s", len(rows), path)
